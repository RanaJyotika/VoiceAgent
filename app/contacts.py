"""Excel import for contacts. Headers are matched to CONTACT_COLUMNS (case-insensitive, with
aliases); every other column is kept as an extra detail of the contact, never dropped."""
import re
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .call_types import CALL_TYPES, DEFAULT, CallType, extra_of
from .db import CONTACT_COLUMNS, IST

# Extra columns the call types read by name (contacts.extra keys), and what sheets call them.
EXTRA_ALIASES = {
    "temple": "temple", "mandir": "temple", "temple_name": "temple",
    "puja_date": "puja_date", "seva_date": "puja_date", "booking_date": "puja_date", "event_date": "puja_date",
    "date": "puja_date", "chadhava_date": "puja_date",
    "devotee_names": "devotee_names", "devotees": "devotee_names", "devotee": "devotee_names",
    "devotee_name": "devotee_names", "sankalp_name": "devotee_names", "sankalp": "devotee_names",
    "puja": "product", "seva": "product", "puja_name": "product", "seva_name": "product", "booking": "product",
    "booking_id": "order_id", "booking_no": "order_id",
    # the template's own headers (call_types Column labels)
    "puja_seva": "product", "puja_or_seva": "product", "payment_cod_prepaid": "payment_mode",
    # abandoned-cart exports (the template's "Items in cart" and "Cart value", and what stores call them)
    "items_in_cart": "product", "cart_items": "product", "cart": "product", "product_names": "product", "line_items": "product",
    "cart_value": "amount", "cart_total": "amount", "cart_amount": "amount", "checkout_value": "amount", "total_price": "amount",
}

ALIASES = {
    "name": "customer_name",
    "customer": "customer_name",
    "mobile": "phone",
    "phone_number": "phone",
    "contact": "phone",
    "item": "product",
    "qty": "quantity",
    "price": "amount",
    "total": "amount",
    "payment": "payment_mode",
    "cod": "payment_mode",
    "delivery": "delivery_date",
    "lang": "language",
    "order": "order_id",
    "order_no": "order_id",
    "order_number": "order_id",
    "customer_name": "customer_name",
    "full_name": "customer_name",
    "phone_no": "phone",
    "mobile_number": "phone",
    "mobile_no": "phone",
    "whatsapp": "phone",
    "product_name": "product",
    "items": "product",
    "quantity": "quantity",
    "amount": "amount",
    "order_amount": "amount",
    "order_total": "amount",
    "payment_mode": "payment_mode",
    "payment_method": "payment_mode",
    "mode": "payment_mode",
    "shipping_address": "address",
    "delivery_address": "address",
    "delivery_date": "delivery_date",
    "expected_delivery": "delivery_date",
    "remarks": "notes",
    "remark": "notes",
    "comment": "notes",
    "comments": "notes",
    "note": "notes",
}


def _normalise(header: str) -> str:
    key = re.sub(r"[^a-z0-9]+", "_", str(header or "").strip().lower()).strip("_")
    return ALIASES.get(key) or EXTRA_ALIASES.get(key) or key


def normalise_phone(raw) -> str:
    """Return +91XXXXXXXXXX style. Accepts 10 digits, 0-prefixed, 91-prefixed, spaces, dashes."""
    text = str(raw or "").strip()
    if text.endswith(".0"):          # Excel turns numbers into floats
        text = text[:-2]
    digits = "".join(ch for ch in text if ch.isdigit())
    if not digits:
        return ""
    if len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    if len(digits) == 10:
        digits = "91" + digits
    return "+" + digits


def is_valid_phone(phone: str) -> bool:
    digits = phone.lstrip("+")
    return phone.startswith("+") and digits.isdigit() and 10 <= len(digits) <= 15


def _cell_text(v) -> str:
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%d %B %Y")    # a date cell
    if isinstance(v, float) and v.is_integer():
        return str(int(v))               # Excel turns 1499 into 1499.0
    return str(v).strip()


def inspect_excel(content: bytes) -> dict:
    """Read the sheet and say what is in it: the rows as contacts, which headers went to the fixed
    columns, which were kept as extra details, and the sheet's own row number for each contact."""
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        return {"rows": [], "columns": [], "extras": []}
    raw_headers = [str(h).strip() if h not in (None, "") else "" for h in header]
    cols = [_normalise(h) if h else "" for h in raw_headers]
    columns = [c for c in cols if c in CONTACT_COLUMNS]
    extras = [(c, h) for c, h in zip(cols, raw_headers) if c and c not in CONTACT_COLUMNS]
    rows: list[dict] = []
    for n, raw in enumerate(rows_iter, start=2):
        if raw is None or all(v in (None, "") for v in raw):
            continue
        row = {c: _cell_text(v) for c, v in zip(cols, raw) if c in CONTACT_COLUMNS}
        row["phone"] = normalise_phone(row.get("phone", ""))
        extra = {c: {"label": h, "value": _cell_text(v)} for (c, h), v in
                 zip([(c, h) for c, h in zip(cols, raw_headers)], raw) if c and c not in CONTACT_COLUMNS and _cell_text(v)}
        row["extra"] = extra
        row["_row"] = n
        rows.append(row)
    return {"rows": rows, "columns": columns, "extras": [{"key": c, "label": h} for c, h in extras]}


def missing_columns(call_type: CallType, columns: list[str], extras: list[dict]) -> list[str]:
    """Required columns of this call type that the sheet does not have (by label)."""
    have = set(columns) | {e["key"] for e in extras}
    return [c.label for c in call_type.columns if c.required and c.key not in have]


def check_rows(rows: list[dict], existing_phones: set[str]) -> list[dict]:
    """Attach a verdict to each row: ok, no_phone, duplicate (same number earlier in the sheet)
    or on_list (same number already on the call list)."""
    seen: set[str] = set()
    out = []
    for r in rows:
        p = r.get("phone", "")
        if not is_valid_phone(p):
            issue = "no_phone"
        elif p in seen:
            issue = "duplicate"
        elif p in existing_phones:
            issue = "on_list"
        else:
            issue = "ok"
        seen.add(p)
        out.append({**r, "_issue": issue})
    return out


def _value(contact: dict, key: str) -> str:
    """A contact's value for a call type column: a fixed column or one of its extras."""
    if key in CONTACT_COLUMNS:
        return str(contact.get(key) or "")
    v = extra_of(contact).get(key)
    return str(v.get("value") if isinstance(v, dict) else v or "")


# What the customer said, in the team's words, by the result the agent recorded.
RESPONSES = {
    "confirmed": "Confirmed", "interested": "Interested", "changed": "Confirmed with changes",
    "cancelled": "Cancelled", "not_ordered": "Says they did not order", "not_interested": "Not interested",
    "already_booked": "Already booked", "needs_help": "Wants help", "call_later": "Call back later",
    "wrong_number": "Wrong number", "do_not_call": "Do not call again", "unclear": "Unclear",
    "wants_to_order": "Wants to order", "already_ordered": "Already ordered",
}
# How a call nobody picked up ended, by live_bridge.classify_unanswered's note (shown as the reason).
NOT_PICKED_UP = {
    "call rejected": "Call rejected", "switched off or not reachable": "Switched off / not reachable",
    "number does not exist": "Number does not exist",
}
DETAIL_LABELS = {
    "not_interested_reason": "Reason", "cancel_reason": "Reason", "callback_time": "Call back",
    "new_delivery_date": "New delivery date", "new_product": "Product they named", "mismatch": "Mismatch",
    "question": "Question", "notes": "Note", "payment_preference": "Payment", "order_changes": "Changes asked for",
}
# A person must check these before acting: the customer changed something, disputed the record, or the
# result is unclear. Owner, 5 Oct 2026: an address the customer dictated always needs a human to verify it.
VERIFY_OUTCOMES = ("unclear", "changed", "not_ordered", "needs_help")
VERIFY_DETAILS = ("new_address", "mismatch", "order_changes", "new_product", "new_delivery_date",
                  "verify")   # set by our own checks, e.g. an address the customer never said yes to (call 135)


def call_result(contact: dict, call: dict | None) -> tuple[str, str, str]:
    """(call status, customer response, why) for one contact, from its latest call. The status is the
    team's next step at a glance: Picked, Not answered, Caller busy, Call after some time or Needs human
    verification (and Not called yet / Calling now / Call dropped / No valid phone number)."""
    from .live_bridge import classify_unanswered   # here: live_bridge pulls in the voice agent
    status = contact.get("status") or "pending"
    if status == "calling":
        return "Calling now", "", ""
    if status == "pending":
        return "Not called yet", "", ""         # never rung, or put back in the queue since its last call
    if not call:
        # A row the campaign could not ring at all is marked wrong_number without a call.
        return ("No valid phone number" if status == "wrong_number" else "Not called yet"), "", ""
    if not call.get("duration_sec"):
        if call.get("error"):
            return "Not called yet", "", ""      # our side failed before it rang; the campaign tries again
        note = classify_unanswered(call.get("hangup_cause") or "")[1]
        if note == "busy":
            return "Caller busy", "", ""
        return "Not answered", "", NOT_PICKED_UP.get(note, "")
    outcome = call.get("outcome") or ""
    if outcome == "dropped":
        return "Call dropped", "", "Picked up, then the call ended before an answer"
    details = call.get("outcome_details") if isinstance(call.get("outcome_details"), dict) else {}
    response = RESPONSES.get(outcome, outcome.replace("_", " ").capitalize())
    if outcome == "call_later":
        return "Call after some time", response, ""
    if outcome in VERIFY_OUTCOMES or any(str(details.get(k) or "").strip() for k in VERIFY_DETAILS):
        return "Needs human verification", response, ""
    return "Picked", response, ""


def build_results(contacts: list[dict], latest_call, call_type: CallType = DEFAULT, campaign_name: str = "") -> bytes:
    """What the team needs from a campaign, one row per contact: the customer's details, why we
    called, whether they picked up, what they said and why, and any address they dictated on the
    call. No transcripts: those stay on the dashboard. latest_call(contact_id) -> the most recent
    call row for that contact, or None."""
    reason = (f"{call_type.name}: {campaign_name}" if campaign_name else call_type.name) if call_type.kind == "promotional" \
        else call_type.name
    header = [c.label for c in call_type.columns] + ["Call reason", "Call status", "Customer response",
                                                     "Reason / details", "Address given on call", "Called on", "Tries"]
    wb = Workbook()
    ws = wb.active
    ws.title = "results"
    ws.append(header)
    for c in contacts:
        call = latest_call(c["id"])
        line, response, why = call_result(c, call)
        details = ((call or {}).get("outcome_details") or {}) if response else {}
        if not isinstance(details, dict):
            details = {}
        said = "; ".join(f"{DETAIL_LABELS.get(k, k.replace('_', ' ').capitalize())}: {details[k]}"
                         for k in call_type.details if k != "new_address" and str(details.get(k) or "").strip())
        when = c.get("last_call_at") or ""
        if when:
            try:
                when = datetime.fromisoformat(when).astimezone(IST).strftime("%d %b %Y %H:%M")
            except ValueError:
                pass
        ws.append([_value(c, col.key) for col in call_type.columns]
                  + [reason, line, response, said or why, str(details.get("new_address") or ""), when, c.get("attempts") or 0])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        widest = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(widest + 2, 50)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template(call_type: CallType = DEFAULT) -> bytes:
    """An empty sheet with this call type's columns and one example row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "contacts"
    ws.append([c.label for c in call_type.columns])
    ws.append([c.example for c in call_type.columns])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def call_type_named(key: str | None) -> CallType:
    return CALL_TYPES.get(str(key or ""), DEFAULT)
